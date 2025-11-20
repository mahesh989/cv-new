"""
Simple CV processor service for basic text extraction
"""
import logging
from pathlib import Path
from typing import Dict, Any, Optional

logger = logging.getLogger(__name__)


class CVProcessor:
    """Simple CV processor for basic text extraction"""
    
    def extract_text_from_file(self, file_path: Path) -> Dict[str, Any]:
        """Extract text from a CV file"""
        print("\n" + "="*80)
        print("🔧 [CV_PROCESSOR] extract_text_from_file() CALLED")
        print("="*80)
        print(f"📁 File path: {file_path}")
        print(f"📄 File exists: {file_path.exists()}")
        print(f"📊 File suffix: {file_path.suffix}")
        print("="*80 + "\n")
        
        try:
            logger.info(f"[CV_PROCESSOR] Processing file: {file_path} (suffix: {file_path.suffix})")
            if file_path.suffix.lower() == '.pdf':
                print("📝 [CV_PROCESSOR] Detected PDF file - using PDF extraction")
                logger.info(f"[CV_PROCESSOR] Using PDF extraction for: {file_path}")
                return self._extract_from_pdf(file_path)
            elif file_path.suffix.lower() == '.docx':
                print("📝 [CV_PROCESSOR] Detected DOCX file - using DOCX extraction")
                logger.info(f"[CV_PROCESSOR] Using DOCX extraction for: {file_path}")
                return self._extract_from_docx(file_path)
            elif file_path.suffix.lower() == '.txt':
                print("📝 [CV_PROCESSOR] Detected TXT file - using TXT extraction")
                logger.info(f"[CV_PROCESSOR] Using TXT extraction for: {file_path}")
                return self._extract_from_txt(file_path)
            else:
                print(f"❌ [CV_PROCESSOR] Unsupported file type: {file_path.suffix}")
                logger.warning(f"[CV_PROCESSOR] Unsupported file type: {file_path.suffix}")
                return {
                    'success': False,
                    'error': f'Unsupported file type: {file_path.suffix}',
                    'text': '',
                    'word_count': 0
                }
        except Exception as e:
            logger.error(f"Error extracting text from {file_path}: {e}")
            return {
                'success': False,
                'error': str(e),
                'text': '',
                'word_count': 0
            }
    
    def _extract_from_pdf(self, file_path: Path) -> Dict[str, Any]:
        """Extract text from PDF file"""
        print("📝 [CV_PROCESSOR] _extract_from_pdf() called")
        try:
            import PyPDF2
            print("✅ [CV_PROCESSOR] PyPDF2 imported successfully")
            with open(file_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                page_count = len(reader.pages)
                print(f"📊 [CV_PROCESSOR] PDF has {page_count} pages")
                text = ""
                for i, page in enumerate(reader.pages):
                    page_text = page.extract_text()
                    text += page_text + "\n"
                    print(f"📝 [CV_PROCESSOR] Extracted page {i+1}/{page_count}: {len(page_text)} chars")
                
                print(f"✅ [CV_PROCESSOR] PDF extraction complete: {len(text)} total characters\n")
                return {
                    'success': True,
                    'text': text.strip(),
                    'word_count': len(text.split()),
                    'file_type': 'pdf'
                }
        except ImportError:
            return {
                'success': False,
                'error': 'PyPDF2 not available',
                'text': '',
                'word_count': 0
            }
    
    def _extract_from_docx(self, file_path: Path) -> Dict[str, Any]:
        """Extract text from DOCX file with bullet point preservation"""
        print("📝 [CV_PROCESSOR] _extract_from_docx() called")
        try:
            from docx import Document
            print("✅ [CV_PROCESSOR] python-docx imported successfully")
            doc = Document(file_path)
            para_count = len(doc.paragraphs)
            print(f"📊 [CV_PROCESSOR] DOCX has {para_count} paragraphs")
            text = ""
            bullet_count = 0
            
            logger.info(f"[DOCX_PROCESSING] Processing DOCX file: {file_path}")
            
            for paragraph in doc.paragraphs:
                para_text = paragraph.text.strip()
                if not para_text:
                    text += "\n"
                    continue
                
                # Enhanced bullet detection for DOCX files
                is_bullet = False
                
                # Pattern 1: Already has bullet symbols
                if para_text.startswith(('•', '-', '*', '◦', '▪', '▫', '→', '►')):
                    is_bullet = True
                
                # Pattern 2: Check for list-style formatting in DOCX
                try:
                    if (paragraph.style.name.startswith('List') or 
                        'Bullet' in paragraph.style.name or
                        'List' in paragraph.style.name):
                        is_bullet = True
                except:
                    pass
                
                # Pattern 2b: Check for numbering properties (bullet/numbered lists)
                try:
                    if paragraph._element.pPr is not None:
                        numPr = paragraph._element.pPr.numPr
                        if numPr is not None:
                            is_bullet = True
                            logger.info(f"[DOCX_PROCESSING] Detected bullet via numbering: {para_text[:50]}...")
                except:
                    pass
                
                # Pattern 3: Detect bullet points by content patterns
                if not is_bullet and len(para_text) > 10 and len(para_text) < 500:
                    # Check if it's not a header (all caps, short)
                    if not (para_text.isupper() and len(para_text) < 50):
                        # Check if it starts with action words (common in bullet points)
                        action_words = [
                            'advanced', 'strong', 'proficient', 'ability', 'excellent', 'developed', 
                            'created', 'implemented', 'managed', 'led', 'designed', 'built', 
                            'analyzed', 'improved', 'reduced', 'increased', 'delivered',
                            'collaborated', 'enhanced', 'optimized', 'automated', 'integrated',
                            'demonstrated', 'applied', 'contributed', 'addressed', 'presented',
                            'technical', 'skilled', 'expertise', 'adept', 'experienced', 'capable'
                        ]
                        if any(para_text.lower().startswith(word) for word in action_words):
                            is_bullet = True
                        
                        # Check for common bullet point patterns
                        bullet_patterns = [
                            'skills', 'experience', 'proficient', 'ability', 'excellent',
                            'reduced', 'improved', 'demonstrated', 'applied', 'contributed',
                            'technical expertise', 'etl processes', 'data analysis', 'data management',
                            'problem-solving', 'adaptability', 'communication', 'team collaboration',
                            'organizational skills', 'skilled in', 'expertise in', 'adept at'
                        ]
                        if any(pattern in para_text.lower() for pattern in bullet_patterns):
                            is_bullet = True
                        
                        # Check for colon-separated bullet points (like "Technical Expertise: ...")
                        if ':' in para_text and len(para_text.split(':')[0]) < 50:
                            is_bullet = True
                
                if is_bullet:
                    bullet_count += 1
                    # Ensure it has a bullet symbol
                    if not para_text.startswith(('•', '-', '*', '◦', '▪', '▫', '→', '►')):
                        text += "• " + para_text + "\n"
                        logger.info(f"[DOCX_PROCESSING] Added bullet: • {para_text[:50]}...")
                    else:
                        text += para_text + "\n"
                        logger.info(f"[DOCX_PROCESSING] Existing bullet: {para_text[:50]}...")
                else:
                    # Regular paragraph
                    text += para_text + "\n"
            
            logger.info(f"[DOCX_PROCESSING] Completed: {bullet_count} bullets detected")
            print(f"✅ [CV_PROCESSOR] DOCX extraction complete: {len(text)} chars, {bullet_count} bullets\n")
            
            return {
                'success': True,
                'text': text.strip(),
                'word_count': len(text.split()),
                'file_type': 'docx'
            }
        except ImportError:
            return {
                'success': False,
                'error': 'python-docx not available',
                'text': '',
                'word_count': 0
            }
    
    def _extract_from_txt(self, file_path: Path) -> Dict[str, Any]:
        """Extract text from TXT file"""
        print("📝 [CV_PROCESSOR] _extract_from_txt() called")
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()
            
            print(f"✅ [CV_PROCESSOR] TXT extraction complete: {len(text)} characters\n")
            return {
                'success': True,
                'text': text.strip(),
                'word_count': len(text.split()),
                'file_type': 'txt'
            }
        except Exception as e:
            return {
                'success': False,
                'error': str(e),
                'text': '',
                'word_count': 0
            }
    
    def get_text_preview(self, text: str, max_length: int = 200) -> str:
        """Get a preview of the text"""
        if len(text) <= max_length:
            return text
        return text[:max_length] + "..."
    
    def extract_basic_info(self, text: str) -> Dict[str, Any]:
        """Extract basic information from CV text"""
        lines = text.split('\n')
        
        # Simple extraction - look for common patterns
        email = ""
        phone = ""
        name = ""
        
        for line in lines:
            line = line.strip()
            if '@' in line and not email:
                email = line
            elif any(char.isdigit() for char in line) and len(line) > 8 and not phone:
                phone = line
            elif len(line) > 2 and len(line) < 50 and not name and not any(char in line for char in ['@', 'http', 'www']):
                name = line
                break
        
        return {
            'name': name,
            'email': email,
            'phone': phone,
            'word_count': len(text.split()),
            'line_count': len(lines)
        }


# Global instance
cv_processor = CVProcessor()
